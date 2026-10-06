# -*- coding: utf-8 -*-
import os, glob, io, subprocess, tempfile
from datetime import datetime
from pptx import Presentation
import fitz
from PIL import Image
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

SOFFICE = r'C:\Program Files\LibreOffice\program\soffice.exe'
SW, SH = 6858000, 9906000
BOX_X, BOX_CX, BOX_CY = 1590474, 4358806, 2403101
PHASES = [('시공전', 1325762), ('시공중', 4278091), ('시공후', 7230419)]
TARGET_W = 300

MASTER = 'TCL_6월 정산/Himart_Logo_Project-시공일자및본사제출_6월분까지.xlsx'
OUT = 'TCL_6월 정산/Himart_Logo_Project-시공일자및본사제출_2~6월_사진완성.xlsx'
MEDIA = 'TCL_6월 정산/_대지크롭_all'
os.makedirs(MEDIA, exist_ok=True)

def point_of(path):
    return os.path.basename(path).replace('TCL_', '').rsplit('.', 1)[0]

# month -> {point: (path, kind)}
month_src = {}
def add(month, paths, kind):
    d = month_src.setdefault(month, {})
    for p in paths:
        d[point_of(p)] = (p, kind)

add(2, [x for x in glob.glob('TCL_2월작업완료 청구서류_pdf/*/TCL_*.pdf') if '견적서' not in x], 'pdf')
add(3, glob.glob('TCL_3월 정산/*/*/TCL_*.pptx'), 'pptx')
add(4, glob.glob('4월 정산/*/견적서/*/TCL_*.pptx'), 'pptx')
add(5, glob.glob('5월 정산/*/견적서/*/TCL_*.pptx'), 'pptx')
add(6, glob.glob('견적서/*/TCL_*.pptx'), 'pptx')

for m in sorted(month_src):
    print(f'  {m}월 대지 인덱스: {len(month_src[m])}개', flush=True)

# batch convert all pptx -> pdf per month into temp dirs
pptx_pdf = {}  # original pptx path -> rendered pdf path
for m, d in month_src.items():
    plist = [p for (p, k) in d.values() if k == 'pptx']
    if not plist:
        continue
    td = tempfile.mkdtemp(prefix=f'm{m}_')
    print(f'  {m}월 PPT->PDF 변환 {len(plist)}개...', flush=True)
    # convert in chunks to avoid command length limits
    CH = 40
    for i in range(0, len(plist), CH):
        chunk = [os.path.abspath(x) for x in plist[i:i+CH]]
        subprocess.run([SOFFICE, '--headless', '--norestore', '--convert-to', 'pdf',
                        '--outdir', td] + chunk,
                       capture_output=True, encoding='utf-8', errors='replace', timeout=600)
    for p in plist:
        cand = os.path.join(td, os.path.splitext(os.path.basename(p))[0] + '.pdf')
        if os.path.exists(cand):
            pptx_pdf[p] = cand

def title_of(path, kind):
    try:
        if kind == 'pptx':
            prs = Presentation(path)
            for sh in prs.slides[0].shapes:
                if sh.has_text_frame:
                    t = ''.join(r.text for pa in sh.text_frame.paragraphs for r in pa.runs)
                    if '지사' in t and '-' in t and '완료사진' not in t:
                        return t.strip()
            for sh in prs.slides[0].shapes:
                if sh.has_text_frame and '지사' in sh.text_frame.text:
                    return sh.text_frame.text.replace('\n', ' ').strip()
        else:
            txt = fitz.open(path)[0].get_text()
            for l in txt.splitlines():
                if '지사' in l and '-' in l:
                    return l.strip()
    except Exception as e:
        return f'(제목추출실패:{e})'
    return '(제목없음)'

def render_crops(pdf_path):
    pix = fitz.open(pdf_path)[0].get_pixmap(dpi=200)
    im = Image.open(io.BytesIO(pix.tobytes('png')))
    W, H = im.size
    out = []
    for nm, by in PHASES:
        c = im.crop((int(BOX_X/SW*W), int(by/SH*H), int((BOX_X+BOX_CX)/SW*W), int((by+BOX_CY)/SH*H)))
        ratio = TARGET_W / c.width
        out.append(c.resize((TARGET_W, int(c.height*ratio))))
    return out

wb = load_workbook(MASTER)
ws = wb['Sheet1']
ws['G1'] = '시공 전'; ws['H1'] = '작업 중'; ws['I1'] = '작업 후'

matched, unmatched, title_fail = [], [], []
COLW = ROWH = None

for r in range(2, 186):
    point = ws.cell(r, 3).value
    dval = ws.cell(r, 2).value
    if not point:
        continue
    point = str(point).strip()
    month = dval.month if isinstance(dval, datetime) else None
    lookup_m = 6 if month == 7 else month
    src = month_src.get(lookup_m, {})
    if point not in src:
        unmatched.append((r, month, point))
        continue
    path, kind = src[point]
    # 3중 검증: 내부 제목에 point 포함?
    title = title_of(path, kind)
    if point.replace(' ', '') not in title.replace(' ', ''):
        title_fail.append((r, point, title))
        continue
    # 렌더 소스 pdf
    pdf = path if kind == 'pdf' else pptx_pdf.get(path)
    if not pdf or not os.path.exists(pdf):
        unmatched.append((r, month, point + '(렌더실패)'))
        continue
    crops = render_crops(pdf)
    for ci, (nm, by) in enumerate(PHASES):
        c = crops[ci]
        fp = os.path.join(MEDIA, f'{lookup_m}_{point}_{nm}.png')
        c.save(fp)
        xi = XLImage(fp); xi.width = c.width; xi.height = c.height
        xi.anchor = f'{chr(71+ci)}{r}'
        ws.add_image(xi)
        COLW, ROWH = c.width, c.height
    ws.row_dimensions[r].height = (ROWH or 165) * 0.75
    matched.append((r, point))

if COLW:
    for col in ('G', 'H', 'I'):
        ws.column_dimensions[col].width = COLW / 7.0

wb.save(OUT)

print('=' * 50, flush=True)
print(f'삽입 완료: {len(matched)}개 지점 × 3 = {len(matched)*3}장', flush=True)
print(f'매칭 실패(폴더에 대지 없음): {len(unmatched)}개', flush=True)
for r, m, p in unmatched:
    print(f'   r{r} [{m}월] {p}', flush=True)
print(f'제목불일치(검증 탈락): {len(title_fail)}개', flush=True)
for r, p, t in title_fail:
    print(f'   r{r} {p} != 제목 {t!r}', flush=True)
print(f'저장: {OUT}', flush=True)

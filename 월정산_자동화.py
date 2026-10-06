# -*- coding: utf-8 -*-
"""
월 정산 자동화 — makeList 시트 하나로 견적서·PPT·총견적서·송부용 PDF까지

사용법 (견적서_자동화 폴더에서):
  _python\\python.exe 월정산_자동화.py make --month 10             # 지점 견적서 + PPT + 총견적서
  _python\\python.exe 월정산_자동화.py pack --month 10             # TCL_10월 정산/ 정리 + PDF + 송부용 zip
  _python\\python.exe 월정산_자동화.py all  --month 10             # make + pack
  옵션: --sheet "10월 정산"  (기본: "{month}월 정산")
        --overwrite          (이미 있는 지점 견적서/PPT도 다시 생성)

흐름:
  makeList.xlsx [N월 정산 시트: 지점명 | 날짜 | 조명유무 | 지사]
    → 견적서/{점}/TCL_{점}.xlsx, .pptx           (견적서_자동화.py 함수 재사용)
    → 견적서/TCL_N월_총견적서.xlsx                (양식/총견적서양식.xlsx 기반, 18줄 넘으면 _1, _2 로 분할)
    → TCL_N월 정산/{지사}/견적서/{점}/            (pack)
    → TCL_N월 정산/{지사}/{지사}_견적서.pdf, _사진.pdf, TCL_N월_총견적서.pdf
    → TCL_N월 정산/N월 정산 송부용/ + .zip

총견적서는 openpyxl 을 쓰지 않고 XML 을 직접 수정한다.
(openpyxl 은 저장 시 병합 셀 바깥 테두리를 첫 칸 테두리로 덮어써서 오른쪽 굵은 선이 깨짐)
"""

import argparse
import copy
import importlib.util
import io
import os
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from lxml import etree

BASE = Path(__file__).resolve().parent
MAKE_LIST = BASE / "makeList.xlsx"
OUTPUT_DIR = BASE / "견적서"
TOTAL_TEMPLATE = BASE / "양식" / "총견적서양식.xlsx"

PRICES = {'조명': 580000, '비조명': 290000}

# 총견적서 양식의 품목 영역: 12~29행 (18줄), 합계 F30
ITEM_FIRST_ROW, ITEM_LAST_ROW = 12, 29
ITEM_ROWS = ITEM_LAST_ROW - ITEM_FIRST_ROW + 1
# 양식에서 지사 행 / 점 행 서식을 가져올 셀 (값은 비어 있음)
JISA_STYLE_CELL, POINT_STYLE_CELL = 'A12', 'A13'

SOFFICE_PATHS = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
]

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def q(tag):
    return f'{{{NS}}}{tag}'


def load_gj():
    """견적서_자동화.py 를 모듈로 불러와 지점 견적서/PPT 생성 함수 재사용"""
    spec = importlib.util.spec_from_file_location('gj', BASE / '견적서_자동화.py')
    gj = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gj)
    return gj


def backup(path: Path):
    """덮어쓰기 전 같은 폴더에 {이름}_백업_{시각}.xlsx 로 복사"""
    if path.exists():
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        dst = path.with_name(f"{path.stem}_백업_{ts}{path.suffix}")
        shutil.copy2(path, dst)
        print(f"  ↩ 백업: {dst.name}")


# ─────────────────────────────────────────────
# makeList 읽기
# ─────────────────────────────────────────────
def read_sheet(sheet_name):
    """반환: [{'name','date','light','jisa'}, ...]  (makeList 순서 유지)"""
    from openpyxl import load_workbook
    wb = load_workbook(MAKE_LIST, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"makeList.xlsx 에 '{sheet_name}' 시트가 없습니다. 시트 목록: {wb.sheetnames}")
    ws = wb[sheet_name]
    header = [str(h).strip() if h else '' for h in next(ws.iter_rows(max_row=1, values_only=True))]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_name, i_date = col('지점명', '지점'), col('날짜')
    i_light, i_jisa = col('조명유무', '조명구분'), col('지사')
    if i_name is None or i_jisa is None:
        raise SystemExit(f"'{sheet_name}' 헤더에 '지점명'과 '지사' 열이 필요합니다. 현재 헤더: {header}")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[i_name]:
            continue
        name = str(r[i_name]).strip()
        name = name if name.endswith('점') else name + '점'
        light = str(r[i_light]).strip() if i_light is not None and r[i_light] else '조명'
        if light not in PRICES:
            light = '조명'
        d = r[i_date] if i_date is not None else None
        rows.append({
            'name': name,
            'date': d if isinstance(d, datetime) else None,
            'light': light,
            'jisa': str(r[i_jisa]).strip() if r[i_jisa] else '',
        })
    missing = [x['name'] for x in rows if not x['jisa']]
    if missing:
        raise SystemExit(f"지사가 비어 있는 지점: {', '.join(missing)}")
    return rows


# ─────────────────────────────────────────────
# 총견적서 (XML 직접 수정)
# ─────────────────────────────────────────────
def _sheet_path(z, sheet_name='견적서'):
    wb = etree.fromstring(z.read('xl/workbook.xml'))
    rid = next(s for s in wb.iter(q('sheet')) if s.get('name') == sheet_name).get(f'{{{R_NS}}}id')
    rels = etree.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    target = next(r for r in rels.iter(f'{{{REL_NS}}}Relationship') if r.get('Id') == rid).get('Target')
    return 'xl/' + target.split('xl/')[-1].lstrip('/')


def _set_cell(c, value, style=None):
    """셀 값 설정: None=비우기, str=인라인 문자열, 숫자=숫자 (서식 s 는 유지/교체)"""
    for ch in list(c):
        c.remove(ch)
    c.attrib.pop('t', None)
    if style is not None:
        c.set('s', str(style))
    if value is None:
        return
    if isinstance(value, str):
        c.set('t', 'inlineStr')
        is_ = etree.SubElement(c, q('is'))
        t = etree.SubElement(is_, q('t'))
        t.text = value
        if value != value.strip():
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    else:
        etree.SubElement(c, q('v')).text = str(value)


def _set_formula_cache(c, value):
    v = c.find(q('v'))
    if v is None:
        v = etree.SubElement(c, q('v'))
    v.text = str(value)


def write_total(out_path: Path, month: int, lines):
    """
    lines: [('jisa', '강원지사'), ('point', '1. 단계점', 580000), ...]  (최대 18줄)
    """
    with zipfile.ZipFile(TOTAL_TEMPLATE) as z_in:
        names = z_in.namelist()
        sheet_xml = _sheet_path(z_in)
        sh = etree.fromstring(z_in.read(sheet_xml))
        cells = {c.get('r'): c for c in sh.iter(q('c'))}

        jisa_style = cells[JISA_STYLE_CELL].get('s')
        point_style = cells[POINT_STYLE_CELL].get('s')

        _set_cell(cells['B4'], f'{month}월')
        _set_cell(cells['B9'], f'TCL 로고 제작설치 작업-{month}월 정산분')

        total = 0
        for i in range(ITEM_ROWS):
            r = ITEM_FIRST_ROW + i
            a, g = cells[f'A{r}'], cells[f'G{r}']
            if i >= len(lines):
                _set_cell(a, None, point_style)
                _set_cell(g, None)
            elif lines[i][0] == 'jisa':
                _set_cell(a, lines[i][1], jisa_style)
                _set_cell(g, '-')
            else:
                _set_cell(a, lines[i][1], point_style)
                _set_cell(g, lines[i][2])
                total += lines[i][2]

        # 수식 캐시값 (재계산 안 하는 뷰어 대비) — F30=SUM, B10=F30, H10=B10
        for ref in ('F30', 'B10', 'H10'):
            _set_formula_cache(cells[ref], total)

        new_sheet = etree.tostring(sh, xml_declaration=True, encoding='UTF-8', standalone=True)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z_out:
            for n in names:
                z_out.writestr(z_in.getinfo(n), new_sheet if n == sheet_xml else z_in.read(n))
    out_path.write_bytes(buf.getvalue())
    return total


def build_total_pages(rows):
    """지사별로 묶어 18줄 단위 페이지로 나눔 (지사 묶음은 가능한 한 한 페이지에, 번호는 페이지마다 1부터)"""
    groups = {}
    for x in rows:
        groups.setdefault(x['jisa'], []).append(x)

    pages, cur = [], []
    for jisa in sorted(groups):
        items = groups[jisa]
        if cur and len(cur) + 1 + len(items) > ITEM_ROWS:   # 지사 묶음이 안 들어가면 다음 장
            pages.append(cur)
            cur = []
        cur.append(('jisa', jisa))
        for x in items:
            if len(cur) >= ITEM_ROWS:                        # 한 지사가 18줄 넘으면 다음 장에 지사명 반복
                pages.append(cur)
                cur = [('jisa', jisa)]
            cur.append(('point', x['name'], PRICES[x['light']]))
    if cur:
        pages.append(cur)

    for page in pages:                                  # 페이지마다 1부터 번호
        n = 0
        for i, ln in enumerate(page):
            if ln[0] == 'point':
                n += 1
                page[i] = ('point', f'{n}. {ln[1]}', ln[2])
    return pages


def make_totals(rows, month):
    if not TOTAL_TEMPLATE.exists():
        raise SystemExit(f"총견적서 양식이 없습니다: {TOTAL_TEMPLATE}")
    pages = build_total_pages(rows)
    outs = []
    for i, page in enumerate(pages, 1):
        suffix = f'_{i}' if len(pages) > 1 else ''
        out = OUTPUT_DIR / f'TCL_{month}월_총견적서{suffix}.xlsx'
        backup(out)
        total = write_total(out, month, page)
        print(f"✅ {out.name}  ({sum(1 for l in page if l[0] == 'point')}개 지점, 합계 {total:,}원)")
        outs.append(out)
    return outs


# ─────────────────────────────────────────────
# make: 지점 견적서 + PPT + 총견적서
# ─────────────────────────────────────────────
def cmd_make(month, sheet, overwrite):
    gj = load_gj()
    rows = read_sheet(sheet)
    print(f"📌 [{sheet}] {len(rows)}개 지점: {', '.join(x['name'] for x in rows)}\n")
    OUTPUT_DIR.mkdir(exist_ok=True)

    for x in rows:
        name, d = x['name'], x['date']
        out_dir = OUTPUT_DIR / name
        out_dir.mkdir(exist_ok=True)

        xlsx = out_dir / f'TCL_{name}.xlsx'
        if xlsx.exists() and not overwrite:
            print(f"⏭  [{name}] 견적서 있음 → 건너뜀")
        else:
            gj._save_excel_zip(gj.get_template_path(x['light']), str(xlsx), {
                'A4': f'{d.year}년' if d else None,
                'B4': f'{d.month}월' if d else None,
                'C4': f'{d.day}일' if d else None,
                'B9': f'TCL 로고 제작설치 작업-{name}',
            })
            print(f"✅ [{name}] 견적서 ({x['light']}, {d:%Y-%m-%d})" if d else f"✅ [{name}] 견적서 (날짜 없음)")

        photo_dir, _ = gj.get_branch_photo_dir(name)
        if photo_dir is None:
            for ph in ('시공전', '시공중', '시공후'):
                (BASE / '사진' / name / ph).mkdir(parents=True, exist_ok=True)
            print(f"   ⚠ [{name}] 사진 폴더 없음 → 사진/{name}/ 생성, PPT 건너뜀")
            continue
        gj.run_make_ppt(name, name, str(out_dir), lambda m: print('  ' + m),
                        allow_overwrite=overwrite, jisa=x['jisa'])

    print()
    make_totals(rows, month)


# ─────────────────────────────────────────────
# pack: 정산 폴더 정리 + PDF 변환/병합 + 송부용 zip
# ─────────────────────────────────────────────
def find_soffice():
    found = shutil.which('soffice') or shutil.which('libreoffice')
    if found:
        return found
    for p in SOFFICE_PATHS:
        if os.path.exists(p):
            return p
    raise SystemExit("LibreOffice(soffice.exe)를 찾을 수 없습니다.")


def to_pdf(soffice, src: Path, dst_pdf: Path, work_root: Path, stem: str, timeout=180):
    """LibreOffice 로 PDF 변환 (한글 경로 문제 피하려고 작업 폴더에 복사 후 변환)"""
    work = work_root / stem
    work.mkdir(parents=True, exist_ok=True)
    copied = work / f"{stem}{src.suffix.lower()}"
    shutil.copy2(src, copied)
    cmd = [soffice, '--headless', '--nologo', '--nofirststartwizard', '--norestore',
           f'-env:UserInstallation={(work / "lo-profile").as_uri()}',
           '--convert-to', 'pdf', '--outdir', str(work), str(copied)]
    res = subprocess.run(cmd, capture_output=True, encoding='utf-8', errors='replace', timeout=timeout)
    produced = work / f"{stem}.pdf"
    if not produced.exists() or produced.stat().st_size == 0:
        raise RuntimeError(f"PDF 변환 실패: {src.name} {(res.stderr or res.stdout or '')[:300]}")
    dst_pdf.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(produced), str(dst_pdf))


def merge_pdfs(inputs, out_pdf: Path):
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for p in inputs:
        for page in PdfReader(str(p)).pages:
            w.add_page(page)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    with out_pdf.open('wb') as fh:
        w.write(fh)
    return len(w.pages)


def safe(s):
    return re.sub(r'[<>:"/\\|?*\s]+', '_', s).strip('_')


def cmd_pack(month, sheet):
    rows = read_sheet(sheet)
    root = BASE / f'TCL_{month}월 정산'
    send = root / f'{month}월 정산 송부용'
    print(f"📦 {root.name} 정리")

    # 1. 원본 복사: 총견적서 + {지사}/견적서/{점}/
    totals = sorted(OUTPUT_DIR.glob(f'TCL_{month}월_총견적서*.xlsx'))
    totals = [t for t in totals if '_백업_' not in t.name]
    if not totals:
        raise SystemExit(f"총견적서가 없습니다. 먼저 make 를 실행하세요: 견적서/TCL_{month}월_총견적서.xlsx")
    root.mkdir(exist_ok=True)
    for t in totals:
        shutil.copy2(t, root / t.name)

    groups = {}
    for x in rows:
        src = OUTPUT_DIR / x['name']
        files = sorted(p for p in src.glob('TCL_*') if p.suffix in ('.xlsx', '.pptx') and not p.name.startswith('~$'))
        if not files:
            raise SystemExit(f"[{x['name']}] 견적서/PPT 가 없습니다: {src}")
        dst = root / x['jisa'] / '견적서' / x['name']
        dst.mkdir(parents=True, exist_ok=True)
        for f in files:
            shutil.copy2(f, dst / f.name)
        if not (dst / f"TCL_{x['name']}.pptx").exists():
            print(f"   ⚠ [{x['name']}] PPT 없음 → 사진 PDF 에서 빠집니다")
        groups.setdefault(x['jisa'], []).append(dst)

    # 2. PDF 변환 / 지사별 병합
    soffice = find_soffice()
    work = root / '_작업'
    shutil.rmtree(work, ignore_errors=True)
    failures = []
    try:
        # 송부용 폴더는 탐색기 등이 잡고 있으면 rmtree 가 실패하므로 파일만 지우고 다시 채움
        send.mkdir(parents=True, exist_ok=True)
        for p in send.rglob('*'):
            if p.is_file():
                p.unlink()
        for t in totals:
            pdf = root / f'{t.stem}.pdf'
            to_pdf(soffice, t, pdf, work, f'000_{safe(t.stem)}')
            shutil.copy2(pdf, send / pdf.name)
            print(f"  ✅ {pdf.name}")

        for jisa in sorted(groups):
            for kind, ext in (('견적서', '.xlsx'), ('사진', '.pptx')):
                done = []
                for idx, d in enumerate(groups[jisa], 1):
                    src = d / f'TCL_{d.name}{ext}'
                    if not src.exists():
                        continue
                    stem = f'{idx:03d}_{safe(jisa)}_{safe(d.name)}_{kind}'
                    out = root / '_개별PDF' / jisa / kind / f'{stem}.pdf'
                    try:
                        to_pdf(soffice, src, out, work, stem)
                        done.append(out)
                    except Exception as e:
                        failures.append(f'{jisa}/{d.name}/{kind}: {e}')
                if done:
                    merged = root / jisa / f'{jisa}_{kind}.pdf'
                    pages = merge_pdfs(done, merged)
                    (send / jisa).mkdir(exist_ok=True)
                    shutil.copy2(merged, send / jisa / merged.name)
                    print(f"  ✅ {jisa}/{merged.name} ({pages}페이지)")
    finally:
        shutil.rmtree(work, ignore_errors=True)

    if failures:
        print("\n❌ 실패 목록")
        for f in failures:
            print('  - ' + f)
        raise SystemExit(1)

    # 3. 송부용 zip
    zpath = root / f'{send.name}.zip'
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        for p in sorted(send.rglob('*')):
            if p.is_file():
                z.write(p, p.relative_to(root))
    print(f"\n🎉 완료: {zpath.relative_to(BASE)} ({zpath.stat().st_size // 1024} KB)")


def main():
    ap = argparse.ArgumentParser(description='TCL 월 정산 자동화')
    ap.add_argument('step', choices=['make', 'pack', 'all'])
    ap.add_argument('--month', type=int, required=True, help='정산 월 (예: 10)')
    ap.add_argument('--sheet', help='makeList 시트 이름 (기본: "{month}월 정산")')
    ap.add_argument('--overwrite', action='store_true', help='지점 견적서/PPT 를 다시 생성')
    a = ap.parse_args()
    sheet = a.sheet or f'{a.month}월 정산'

    if a.step in ('make', 'all'):
        cmd_make(a.month, sheet, a.overwrite)
    if a.step in ('pack', 'all'):
        print()
        cmd_pack(a.month, sheet)


if __name__ == '__main__':
    main()
